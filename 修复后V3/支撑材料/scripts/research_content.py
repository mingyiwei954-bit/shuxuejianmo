"""Inject measured diagnostics and a byte-verifiable complete source appendix."""
from pathlib import Path
import json, hashlib
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from paper_content import table,fmt
ROOT=Path(__file__).resolve().parents[1]
FINAL=ROOT/'outputs/c_final_v1'
R=FINAL/'research'

def source_files():
    files=[]
    for folder in ['src','scripts','tests']:
        files.extend(p for p in (ROOT/folder).rglob('*') if p.is_file() and p.suffix in ['.py','.mjs'] and '__pycache__' not in p.parts)
    files.extend([ROOT/'run_pipeline.sh',ROOT/'requirements.txt',ROOT/'config/final.yaml'])
    return sorted(files,key=lambda p:str(p.relative_to(ROOT)))

def append_sources(path):
    doc=Document(path)
    h=doc.add_paragraph('附录C 完整源程序与运行配置',style='Heading 2');h.paragraph_format.page_break_before=True
    doc.add_paragraph('以下逐文件收录当前工程自编的全部 Python、JavaScript、Shell 程序和配置。第三方依赖按 requirements.txt 安装；工作簿导出还需 Node 与 @oai/artifact-tool，文档转换需配套 LibreOffice 运行环境。程序以支撑材料内的原始文件运行，版面换行不改变源文件。独立数值复验仅依赖 Python，不要求文档转换环境。')
    records=[]
    for p in source_files():
        name=str(p.relative_to(ROOT));raw=p.read_bytes();content=raw.decode('utf-8')
        doc.add_paragraph(name,style='Heading 3')
        note=doc.add_paragraph('SHA256 '+hashlib.sha256(raw).hexdigest())
        note.paragraph_format.space_after=Pt(2)
        for run in note.runs:run.font.size=Pt(7)
        for line in content.splitlines():
            para=doc.add_paragraph();para.paragraph_format.space_after=Pt(0);para.paragraph_format.space_before=Pt(0)
            para.paragraph_format.line_spacing=1.0;para.paragraph_format.keep_together=False
            run=para.add_run(line if line else ' ');run.font.name='Consolas';run.font.size=Pt(7.5)
            run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
        records.append({'path':name,'sha256':hashlib.sha256(raw).hexdigest(),'lines':len(content.splitlines())})
    doc.save(path)
    (FINAL/'evidence/source_appendix_manifest.json').write_text(json.dumps(records,ensure_ascii=False,indent=2))

def enrich(text,ai):
    annual=pd.read_csv(R/'annual_solver_comparison.csv');tail=pd.read_csv(R/'tail_risk.csv');err=pd.read_csv(R/'prediction_error.csv');blocks=pd.read_csv(R/'block_sensitivity.csv');months=pd.read_csv(R/'monthly_mechanisms.csv')
    start=text.index('针对负载');end=text.index('\n\n关键词',start)
    text=text[:start]+'''针对微网负载、光伏与电价不确定性，本文以十分钟物理区间为统一尺度，将储能调度、非对称缺口损失与购电承诺建立为可实施的联合优化模型。数据仅来自题目附件，所有预测输入满足观测完成和预报发布时间约束。

问题一采用含效率、功率与储电量边界的线性规划，得到日购电费35,126.95元。问题二和问题四的日前模型把五倍价紧急采购直接纳入经验场景目标；历史负载、光伏和电价按同日联合取样，候选场景与终端惩罚只在一月选定。二月至十二月费用分别为14,794,949.72元和15,623,134.83元，相对均值预测控制分别减少1,662,260.45元和1,668,740.06元。

问题三及其波动电价扩展采用24小时滚动窗口。固定共同日前承诺后，0、6、12、18点更新的S3方案相对仅0点更新的S0方案分别减少123,506.61元和155,342.73元；同时报告S2、不同求解器和计费解释的结果，不把单一路径排序推广为信息价值定理。对真实承诺链逐笔结算，逐次调整相对最终净额分别多96,160.53元和98,826.62元。

为避免无历史首日的信息泄漏，一月采用明示的保守启动合同，并从4464个实际执行时段衔接二月SOC。以独立物理复算、共同基线、全年替代求解器回放、月度机制分析及块长度敏感性检验形成证据链。结果支持在当前数据与计费解释下显式处理尾部缺口损失和日内信息更新；单个有限窗口的最优解不等于全年策略最优，经验尾部统计也不能保证未知极端场景下的供电表现。'''+text[end:]
    text=text.replace('微网由小区负载、光伏电源、储能和外部电网构成。','题目[1]中的微网由小区负载、光伏电源、储能和外部电网构成。',1)
    intro='''

### 1.1 方法依据与本题贡献

线性规划提供物理可行域和费用目标的求解基础[2]；样本平均近似以有限场景表示不确定量[3]；滚动控制在新观测到达后更新有限窗口决策[4]。这些是已有方法，本文不把算法名称本身列为创新。本题的针对性工作是：将五倍缺口代价、联合历史误差和储能跨时段耦合放入同一规划；以共同日前合同隔离日内更新频次；记录每次承诺变化，使净额与逐次结算可在同一执行路径上核验。

本文分别回答三个问题：给定信息下决策是否可实施，计费与物理结果是否可复算，以及相对可行基线是否改善经济结果。前两者不能替代第三者；对有限窗口的求解最优性也不作全年策略最优性解释。

### 1.2 技术路线

原始附件先按日期和物理区间对齐并统一量纲；每次决策截取已完成历史和已发布预报，生成联合场景；在SOC和承诺约束下求解，仅执行至下一允许更新时点；区间结束后记录实际缺口、费用和下一状态。训练期选择、正式回放、对照实验及文稿生成分别留存输出，避免把预测样本当成实测数据。本文使用历史数据回放，因此因果信息检查证明的是逐次决策的信息边界，不宣称真实部署过该系统。
'''
    text=text.replace('\n## 2 假设与符号',intro+'\n## 2 假设与符号')
    proof='''

上述互斥处理具有支配关系。若z≥0，则C′=C−D/(ηcηd)、D′=0，减少的母线需求为D(1/(ηcηd)−1)≥0；若z<0，则C′=0、D′=D−ηcηdC，减少的母线需求为C(1−ηcηd)≥0。两种情形均有C′≤C、D′≤D，故额定功率上界保留，SOC轨迹不变，采购需求不增。由于本题允许不用富余且不设最小充放功率，此变换给出不劣的可行互斥解。若另加禁止弃电、负电价收益或最小运行时间等条件，需要重新论证，不能直接沿用该结论。
'''
    text=text.replace('\n问题一以Q=0',proof+'\n问题一以Q=0')
    text=text.replace('每一历史误差也必须相对当时仅用更早日期生成的预测计算，不能用当前的拟合曲线回填过去。','每一历史误差也必须相对当时仅用更早日期生成的预测计算，不能用当前的拟合曲线回填过去。这遵循逐次向前验证的时间切分原则[5]。')
    text=text.replace('本文使用完整LP进行正式调度。','本文使用完整LP进行正式调度。更具体地，单期目标对采购q的导数为p−5p·Pr(N>q)，令其为零得到0.8分位点；离散经验分布取对应分位区间即可。若价格和缺口相关，期望应对联合分布计算，不能先把价格与需求独立平均。')
    text=text.replace('\n## 6 连续启动与因果检验','''

### 5.3 信息价值与实际收益的区别

如果扩大的信息集允许完全忽略新增预报，且优化覆盖相同状态合同下的全部可行策略，则较丰富信息的最优期望费用不会更高：旧策略仍属于新策略集合。这是可行策略集合包含关系的结果，而不是某次实现路径上的费用保证。本文采用有限窗口、有限经验场景和数值求解，因此既不保证每次更新降低实际费用，也不以S2略低于S3否定信息的一般价值。是否引入更新应依据受控费用、缺口和实现负担共同判断。

## 6 连续启动与因果检验''')
    text=text.replace('逐日配对差使用7天移动块','逐日配对差参考块自助法处理局部时间相关的思路[6]，使用7天移动块')
    text=text.replace('分别用HiGHS自动、对偶单纯形和内点法复求。','分别用HiGHS自动、对偶单纯形和内点法复求。HiGHS求解器及其对偶单纯形实现的来源见文献[7]。')
    text=text.replace('全年策略排序仍需额外完整回放才能确认。','其全年影响进一步通过第7.3节的连续回放检验。')
    rows=[]
    for (mode,method),g in annual.groupby(['mode','method'],sort=False):
        a=g.set_index('scenario').cost;rows.append([mode,method,fmt(a['S2']),fmt(a['S3']),fmt(a['S3']-a['S2'])])
    diagnostics='\n\n### 7.1 缺口与尾部成本\n\n'
    diagnostics+='紧急购电补足实际缺口，以下指标衡量额外采购暴露，不等同于停电概率。每日费用最高的17天定义为本样本上尾约5%，其均值仅作描述性尾部指标；本文没有把目标函数改为CVaR。\n\n'
    diagnostics+=table(['模式','紧急时段占比','最高日费/元','上尾17日均费/元','紧急费占比'],[[r['mode'],fmt(r.emergency_interval_fraction*100)+'%',fmt(r.max_daily_cost),fmt(r.tail5_mean_cost),fmt(r.emergency_cost_share*100)+'%'] for _,r in tail.iterrows()])
    diagnostics+='\n\n### 7.2 预测误差与月度机制\n\n预测误差针对实际执行时段使用的场景均值，Q3包含多个发布时刻，不能将其与Q2的差异单独归因为预测模型。MAE=mean(|预测−实测|)，RMSE=sqrt(mean((预测−实测)²))；夜间光伏和接近零的量不适合用MAPE。原始观测未被改动。\n\n'
    diagnostics+=table(['模式','变量','MAE','RMSE','偏差'],[[r['mode'],{'load':'负载/kWh','pv':'光伏/kWh','price':'电价/元每kWh'}[r.variable],fmt(r.mae,4),fmt(r.rmse,4),fmt(r.bias,4)] for _,r in err.iterrows() if r.variable!='price' or r['mode'].startswith('q4')])
    diagnostics+='\n\n月度统计见图5：成本波动同时取决于负荷、电价和充放电轨迹，不能仅由一个误差指标解释。\n\n[[FIGURE:monthly_risk.png|正式期月度费用与紧急采购占比]]\n\n'
    for mode in ['q2','q4_2','q3','q4_3']:
        z=months[months['mode']==mode];a=z.loc[z.cost.idxmax()];b=z.loc[(z.emergency_cost/z.cost).idxmax()]
        diagnostics+=f'{mode}月总费最高为{a.month}，{fmt(a.cost)}元；紧急费用占比最高为{b.month}，{fmt(b.emergency_cost/b.cost*100)}%。\n\n'
    diagnostics+='### 7.3 全年替代求解器回放\n\n在补充回放前固定协议，以相同一月末SOC、同一组S0参考日前承诺和一月已选定参数，分别对两种电价的S2、S3运行HiGHS内点法。每组连续回放334天，不在每日重置状态；保留每段动作并从原始附件重算费用。四组结果全部报告，不据本实验更换正式方案。\n\n'+table(['模式','求解方法','S2费用/元','S3费用/元','S3减S2/元'],rows)
    for mode in ['q3','q4_3']:
        g=annual[annual['mode']==mode].pivot(index='method',columns='scenario',values='cost');delta=g.S3-g.S2
        diagnostics+=f'\n\n{mode}中，默认方法的S3减S2为{fmt(delta["highs"])}元，内点法为{fmt(delta["highs-ipm"])}元。'+('两种方法的排序方向一致，但差距大小随数值方法变化。' if np.sign(delta['highs'])==np.sign(delta['highs-ipm']) else '两种方法出现排序反转，不能保留S2严格优于S3的稳健性结论。')
    diagnostics+='\n\n### 7.4 时间相关与块长度敏感性\n\n在同一334天费用差上，预先列出3、7、14、28天四种移动块长度，每种重采样2000次，使用固定随机种子20260913。表中列S3减S2，正数表示S3较贵；完整S3减S0区间也随支撑材料交付。年份存在季节变化，块重采样不保证严格平稳条件，区间用于判断样本内结论对重采样尺度的敏感性，不解释为下一年收益保证。\n\n'
    z=blocks[blocks.comparison=='S3-S2'];diagnostics+=table(['模式','块长/天','实际差额/元','95%下端/元','95%上端/元'],[[r['mode'],int(r.block_days),fmt(r.difference),fmt(r.ci_low),fmt(r.ci_high)] for _,r in z.iterrows()])
    diagnostics+='\n\n四种块长度下，S3相对S0的费用差区间均低于零，S3相对S2的区间均跨零。日内更新具有本样本经济价值，但尚不足以认定增加18点更新严格改善经济性。将更新计算或交易操作成本加入后，还需重新计算其盈亏阈值。\n\n'
    text=text.replace('\n## 8 证据链与结论',diagnostics+'\n## 8 证据链与结论')
    text=text.replace('文档构建从持久化paper/main.md读取内容，全部结果宏由当前运行值替换；生成后再核对指定区间和四小时聚合。','文档构建从持久化正文与补充诊断生成器读取内容，全部表格由对应冻结值生成；生成后再核对指定区间和四小时聚合。')
    a=text.index('## 参考材料');b=text.index('## 附录',a)
    text=text[:a]+'''## 参考文献

[1] 全国大学生数学建模竞赛组委会. 2026年高教社杯全国大学生数学建模竞赛C题：微网与外部电网电力调控策略及附件[Z]. 2026.

[2] BOYD S, VANDENBERGHE L. Convex Optimization[M]. Cambridge: Cambridge University Press, 2004. https://web.stanford.edu/~boyd/cvxbook/.

[3] SHAPIRO A, DENTCHEVA D, RUSZCZYNSKI A. Lectures on Stochastic Programming: Modeling and Theory[M]. 3rd ed. Philadelphia: SIAM, 2021. DOI:10.1137/1.9781611976595.

[4] RAWLINGS J B, MAYNE D Q, DIEHL M M. Model Predictive Control: Theory, Computation, and Design[M]. 2nd ed. Santa Barbara: Nob Hill Publishing, 2017. https://sites.engineering.ucsb.edu/~jbraw/mpc/.

[5] HYNDMAN R J, ATHANASOPOULOS G. Forecasting: Principles and Practice[M/OL]. 3rd ed. Melbourne: OTexts, 2021[2026-09-12]. https://otexts.com/fpp3/tscv.html.

[6] KUNSCH H R. The jackknife and the bootstrap for general stationary observations[J]. The Annals of Statistics, 1989, 17(3): 1217-1241. DOI:10.1214/aos/1176347265.

[7] HUANGFU Q, HALL J A J. Parallelizing the dual revised simplex method[J]. Mathematical Programming Computation, 2018, 10(1): 119-142. DOI:10.1007/s12532-017-0130-5.

'''+text[b:]
    files=source_files()
    text+='\n\n支撑材料的必要文件按功能列示如下；每个源码文件在附录C逐一给出，完整文件哈希清单见包内SHA256SUMS.json。\n\n'+table(['目录或文件','内容'],[['src/model','全部模型与数据加载程序'],['src/audit','独立逐时段及工作簿审核程序'],['scripts','运行、补充实验、图表、文稿、打包程序'],['tests','数学约束和未来观测扰动测试'],['config/final.yaml 与 requirements.txt','配置及依赖'],['frozen 与 workbooks','原正式结果及五份指定工作簿'],['research','补充全年回放、尾部、月度、误差与统计诊断'],['evidence 与 audit','来源、源码附录清单及复验记录'],['charts 与 logs','图像及运行日志'],['AI工具使用详情.pdf','真实使用范围及仍需人工核验事项']])
    ai+='''

## 本轮继续修订的真实提示与采用情况

用户要求“我同意 我们要解决问题 不用假的数据 不用未来函数 建模不做成通俗化建模 注意各种数学建模的约束 开始做”，后续要求“能继续做吗 继续做呗 给过你链路 记得我要高分啊”。本轮据此恢复已交付快照，补做两种电价下S2和S3的全年内点法回放，独立复算其物理与费用指标；补充月度、尾部、预测误差及块长度诊断，修订论文和完整源码附录。

提供的两份自查DOCX及速成PDF仅作为参考材料。未采用其中示例数值、固定评分权重和获奖保证；适用的修改方向包括有数据支持的检验、文献对应和完整代码。学术文献由作者或出版社公开页面核对，文稿不声称这些既有方法由本队原创。

## 人工核验状态

现有可核实人工行为为用户提出约束、审阅对话结论并授权修订和仓库同步。核心代码修订、实验组织和本轮正文补充由AI实际参与完成，未降格描述为仅改错别字。尚无证据表明参赛队员已经逐项独立复核模型推导、程序和结果，因此这些项登记为待核验；不能由AI自检代替人工主导，也不能以本说明冒充团队完成签核。此前会话完整提示及采用修改记录存在缺口，参赛队需据真实记录补齐。
'''
    return text,ai
