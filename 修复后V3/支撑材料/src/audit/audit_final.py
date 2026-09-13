"""Independent validation from original workbooks and public dispatch/commitment data.
Does not import the model, scenario generator or model settlement functions.
"""
from pathlib import Path
from datetime import datetime
import json,csv,argparse,hashlib,os
import numpy as np
import pandas as pd
import openpyxl
from scipy.optimize import linprog
ROOT=Path(__file__).resolve().parents[2]

def run(workbooks=False):
    cfg=json.loads((ROOT/'config/final.yaml').read_text());F=ROOT/cfg['output_root']/'frozen';O=F.parent/'audit';O.mkdir(parents=True,exist_ok=True);checks=[];metrics={}
    def check(name,ok,actual):checks.append(dict(check=name,status='PASS' if bool(ok) else 'FAIL',actual=actual))
    P=Path(os.environ.get('CUMCM_SOURCE_ROOT',cfg['source_root']))/'C题/附件'
    def matrix(path,sheet):
        w=openpyxl.load_workbook(path,read_only=True,data_only=True);rows=list(w[sheet].values);w.close();return np.array(rows[1:],object)[:,1:145].astype(float)
    w=openpyxl.load_workbook(P/'附件1.xlsx',read_only=True,data_only=True);a=np.array(list(w.active.values)[1:],object);w.close();price=a[:,1].astype(float);qL=a[:,2].astype(float)/6;qV=a[:,3].astype(float)/6
    L=matrix(P/'附件2.xlsx','小区负载')/6;V=matrix(P/'附件2.xlsx','光伏发电实际功率')/6;price_rt=matrix(P/'附件4.xlsx','Sheet1');models={};tol=1e-7
    for path in sorted(F.glob('*dispatch.csv')):
        if 'january' in path.name or path.name.startswith('research_'):continue
        d=pd.read_csv(path);name=path.stem;mode='q1' if name.startswith('q1') else ('q4_3' if name.startswith('q4_3') else 'q4_2' if name.startswith('q4_2') else name[:2]);models[name]=d
        if mode=='q1':l,v,p=qL,qV,price
        else:
            i=(pd.to_datetime(d.date)-pd.Timestamp('2025-01-01')).dt.days.to_numpy();j=d.position.to_numpy()-1;l=L[i,j];v=V[i,j];p=price_rt[i,j] if mode.startswith('q4') else price[j]
        c=d.charge_kwh.to_numpy();dis=d.discharge_kwh.to_numpy();g=d.final_adjusted_grid_kwh.to_numpy();b=d.plan_00_grid_kwh.to_numpy();s0=d.soc_start_kwh.to_numpy();s1=d.soc_end_kwh.to_numpy();q=np.maximum(l+c-v-dis-g,0);u=np.maximum(g+v+dis-l-c,0);delta=g-b
        cost=p*b+p*(delta+.5*np.abs(delta))+5*p*q
        maxerr=lambda x:float(np.max(np.abs(x)))
        check(name+'.count',len(d)==(144 if mode=='q1' else 48096) and not d.duplicated(['date','position']).any() and all(len(z)==144 and set(z.position)==set(range(1,145)) for _,z in d.groupby('date')),len(d))
        check(name+'.raw_load_pv_price',max(maxerr(l-d.actual_load_kwh),maxerr(v-d.actual_pv_kwh),maxerr(p-d.settlement_price))<tol,dict(load=maxerr(l-d.actual_load_kwh),pv=maxerr(v-d.actual_pv_kwh),price=maxerr(p-d.settlement_price)))
        labels=[f'{(int(x)-1)*10//60:02d}:{(int(x)-1)*10%60:02d}-{int(x)*10//60:02d}:{int(x)*10%60:02d}' for x in d.position]
        check(name+'.physical_interval',labels==d.physical_interval.tolist(),sum(x!=y for x,y in zip(labels,d.physical_interval)))
        sr=maxerr(s1-s0-.9*c+dis/.9);continuity=maxerr(s0[1:]-s1[:-1]);check(name+'.soc',sr<tol and continuity<tol,dict(recursion=sr,continuity=continuity))
        check(name+'.device_constraints',min(s0.min(),s1.min())>=1200-tol and max(s0.max(),s1.max())<=10800+tol and c.min()>=-tol and dis.min()>=-tol and max(c.max(),dis.max())<=5000/6+tol and not np.any((c>tol)&(dis>tol)),dict(min_soc=float(min(s0.min(),s1.min())),max_soc=float(max(s0.max(),s1.max())),max_charge=float(c.max()),max_discharge=float(dis.max())))
        check(name+'.realized_balance',max(maxerr(q-d.emergency_kwh),maxerr(u-d.unused_surplus_kwh))<tol,maxerr(q-d.emergency_kwh))
        check(name+'.main_cost',maxerr(cost-d.final_relative_cost_yuan)<tol,maxerr(cost-d.final_relative_cost_yuan))
        check(name+'.same_path_inequality',float((d.sequential_cost_yuan-d.final_relative_cost_yuan).min())>=-tol,float((d.sequential_cost_yuan-d.final_relative_cost_yuan).min()))
        check(name+'.end_soc',abs(s1[-1]-6000)<tol,float(s1[-1]))
        if mode!='q1':
            jan=pd.read_csv(F/f'{mode}_january_dispatch.csv');check(name+'.January_link',abs(s0[0]-jan.soc_end_kwh.iloc[-1])<tol,dict(january_end=float(jan.soc_end_kwh.iloc[-1]),february_start=float(s0[0])))
        metrics[name]=dict(total_cost=float(cost.sum()),plan_cost=float((p*b).sum()),adjustment_cost=float((p*(delta+.5*abs(delta))).sum()),emergency_cost=float((5*p*q).sum()),emergency_kwh=float(q.sum()),emergency_intervals=int((q>=.00005).sum()),max_emergency_kwh=float(q.max()))
    summary=pd.read_csv(F/'experiment_summary.csv')
    for r in summary.itertuples():
        key=(r.mode+'_dispatch') if r.mode in ['q2','q4_2'] and r.model_arm=='joint_empirical' else (r.mode+'_'+r.model_arm+'_dispatch' if r.mode in ['q2','q4_2'] else r.mode+'_'+r.scenario+'_dispatch')
        m=metrics[key];err=max(abs(m['total_cost']-r.final_relative_cost_yuan),abs(m['emergency_kwh']-r.emergency_kwh),abs(m['plan_cost']-r.plan_cost_yuan),abs(m['adjustment_cost']-r.adjustment_cost_yuan),abs(m['emergency_cost']-r.emergency_cost_yuan))
        check(key+'.summary_reconciliation',err<1e-6,err)
    # Reconstruct every saved schedule's formal chain, including repeated revisions.
    ledgers=[]
    for mode in ['q3','q4_3']:
        base=models[f'{mode}_S0_dispatch'].plan_00_grid_kwh.to_numpy()
        for sc in ['S0','S1','S2','S3']:
            d=models[f'{mode}_{sc}_dispatch'];check(f'{mode}.{sc}.common_plan',np.max(np.abs(base-d.plan_00_grid_kwh))<tol,float(np.max(np.abs(base-d.plan_00_grid_kwh))))
            dd=d.set_index(['date','position']);v=pd.read_csv(F/f'{mode}_{sc}_decisions.csv',low_memory=False);chain={};fees={};bad=0;past=0
            prices={(r.date,int(r.position)):float(r.settlement_price) for r in d.itertuples()}
            for r in v.itertuples():
                if r.is_preview_next_day:continue
                k=(r.target_date,int(r.position));new=float(r.new_commitment_kwh)
                delivery=pd.Timestamp(r.target_date)+pd.Timedelta(minutes=(int(r.position)-1)*10)
                past+=int(delivery<pd.Timestamp(r.decision_time))
                if k in chain:
                    bad+=int(abs(chain[k]-r.old_commitment_kwh)>tol);delta=new-chain[k];fees[k]=fees.get(k,0)+prices[k]*(delta+.5*abs(delta))
                chain[k]=new
            last=np.array([chain[k] for k in dd.index]);computed=dd.plan_cost_yuan.to_numpy()+dd.emergency_cost_yuan.to_numpy()+np.array([fees.get(k,0) for k in dd.index]);err=float(np.max(abs(computed-dd.sequential_cost_yuan.to_numpy())))
            check(f'{mode}.{sc}.formal_chain',bad==0 and past==0 and np.max(abs(last-dd.final_adjusted_grid_kwh.to_numpy()))<tol and err<tol,dict(old_commitment_errors=bad,past_rewrites=past,final_error=float(np.max(abs(last-dd.final_adjusted_grid_kwh.to_numpy()))),sequential_error=err))
            ledgers.append(dict(mode=mode,scenario=sc,formal_chain_cost_yuan=float(computed.sum()),main_cost_yuan=float(dd.final_relative_cost_yuan.sum()),difference_yuan=float(computed.sum()-dd.final_relative_cost_yuan.sum())))
    for mode in ['q2','q4_2','q3','q4_3']:
        d=pd.read_csv(F/f'{mode}_january_dispatch.csv');i=(pd.to_datetime(d.date)-pd.Timestamp('2025-01-01')).dt.days.to_numpy();j=d.position.to_numpy()-1;p=price_rt[i,j] if mode.startswith('q4') else price[j];q=np.maximum(L[i,j]-V[i,j]-d.plan_grid_kwh,0);cost=p*d.plan_grid_kwh+5*p*q
        history_ok=all(all(h<r.date for h in str(r.history_dates).split('|') if h and h!='nan') for r in d.itertuples())
        check(mode+'.January_startup',len(d)==4464 and history_ok and np.all(d.charge_kwh==0) and np.all(d.discharge_kwh==0) and np.max(abs(d.soc_start_kwh-6000))<tol and np.max(abs(d.soc_end_kwh-6000))<tol and np.max(abs(cost-d.cost_yuan))<tol,dict(rows=len(d),history_before_decision=history_ok,cost_error=float(np.max(abs(cost-d.cost_yuan))),cost=float(cost.sum())))
    stats=pd.read_csv(F/'solver_decisions.csv');bad=int((pd.to_datetime(stats.latest_observation_time)>pd.to_datetime(stats.decision_time)).sum());check('observation_availability',bad==0,bad)
    # Independent Q1 LP with raw input and physical units, separate construction.
    n=144;M=np.zeros((2*n,5*n));rhs=np.r_[qL-qV,np.zeros(n)];rhs[n]=6000
    for t in range(n):
        M[t,t]=1;M[t,n+t]=-1;M[t,2*n+t]=1;M[t,3*n+t]=-1;M[n+t,n+t]=-.9;M[n+t,2*n+t]=1/.9;M[n+t,4*n+t]=1
        if t:M[n+t,4*n+t-1]=-1
    bounds=[(0,None)]*n+[(0,5000/6)]*(2*n)+[(0,x) for x in qV]+[(1200,10800)]*n;bounds[-1]=(6000,6000)
    opt=linprog(np.r_[price,np.zeros(4*n)],A_eq=M,b_eq=rhs,bounds=bounds,method='highs');check('Q1.independent_optimum',opt.success and abs(opt.fun-metrics['q1_dispatch']['total_cost'])<1e-6,float(opt.fun))
    if workbooks:
        for mode,name in [('q1','result1'),('q2','result2'),('q3','result3'),('q4_2','result4-2'),('q4_3','result4-3')]:
            d=models[mode+'_dispatch'];w=openpyxl.load_workbook(F.parent/f'workbooks/{name}.xlsx',read_only=True,data_only=True)
            required={'计划购电量','充放电量','指定时段','时段映射'}
            if mode!='q1':required.add('紧急购电量')
            if mode in ['q3','q4_3']:required.add('调整购电量')
            check(name+'.required_sheets',required<=set(w.sheetnames),sorted(w.sheetnames))
            for sheet,col in [('计划购电量','plan_00_grid_kwh'),('调整购电量','final_adjusted_grid_kwh')]:
                if sheet not in w.sheetnames:continue
                rows=list(w[sheet].values)
                vals=np.array([r[1] for r in rows[1:145]],float) if mode=='q1' else np.array([r[1:145] for r in rows[1:335]],float).reshape(-1)
                e=float(np.max(abs(vals-d[col].to_numpy())));check(f'{name}.{sheet}',e<=.000051,e)
                if mode!='q1':
                    cost_col='plan_cost_yuan' if sheet=='计划购电量' and mode in ['q3','q4_3'] else 'final_relative_cost_yuan'
                    expected=d.groupby('date',sort=False)[[col,cost_col]].sum().to_numpy()
                    actual=np.array([r[145:147] for r in rows[1:335]],float)
                    err=float(np.max(abs(expected-actual)))
                    dates_ok=[str(r[0])[:10] for r in rows[1:335]]==list(d.date.unique())
                    check(f'{name}.{sheet}.daily_totals_and_dates',err<=.000051 and dates_ok,dict(max_error=err,dates_match=dates_ok))
            # All storage blocks and SOC endpoints; all emergency interval keys.
            dates=d.date.unique();srows=list(w['充放电量'].values)[1:];errs=[]
            for di,day in enumerate(dates):
                z=d[d.date==day]
                for b in range(6):
                    rr=srows[di*6+b];off=1 if mode=='q1' else 2;zz=z.iloc[b*24:(b+1)*24]
                    errs.extend([abs(float(rr[off])-zz.charge_kwh.sum()),abs(float(rr[off+1])-zz.discharge_kwh.sum())])
                    if b<2:errs.append(abs(float(rr[off+3])-(z.soc_start_kwh.iloc[0] if b==0 else z.soc_end_kwh.iloc[-1])))
            check(f'{name}.all_storage_blocks',max(errs)<=.000051,float(max(errs)))
            if mode!='q1':
                expected=[]
                for day,z in d.groupby('date'):
                    active=[]
                    def emit():
                        a=(active[0].position-1)*10;b=active[-1].position*10
                        expected.append((day,f'{a//60}:{a%60:02d}-{b//60}:{b%60:02d}',sum(x.emergency_kwh for x in active)))
                    for r in z.itertuples():
                        if r.emergency_kwh>=.00005:active.append(r)
                        elif active:emit();active=[]
                    if active:emit()
                actual=[];day=None
                for r in list(w['紧急购电量'].values)[1:]:
                    if r[0] is not None:day=str(r[0])[:10]
                    actual.append((day,r[1],float(r[2])))
                keys_ok=[x[:2] for x in actual]==[x[:2] for x in expected]
                err=max([abs(x[2]-y[2]) for x,y in zip(actual,expected)] or [0])
                check(name+'.emergency_intervals',keys_ok and err<=.000051,dict(rows=len(actual),expected=len(expected),value_error=err))
            specified=list(w['指定时段'].values)[1:];specified_error=[]
            for rr in specified:
                day=str(rr[0])[:10];zz=d[(d.date==day)&(d.physical_interval==rr[1])]
                if len(zz)!=1:specified_error.append('missing key')
                else:
                    x=zz.iloc[0]
                    if abs(float(rr[2])-x.plan_00_grid_kwh)>.000051 or abs(float(rr[3])-x.final_adjusted_grid_kwh)>.000051:specified_error.append((day,rr[1]))
            check(name+'.specified_physical_values',not specified_error and len(specified)==(6 if mode=='q1' else 24),specified_error)
            w.close()
    out=dict(status='PASS' if all(r['status']=='PASS' for r in checks) else 'FAIL',generated_at=datetime.now().astimezone().isoformat(),checks=checks,metrics=metrics,formal_settlement=ledgers)
    (O/'independent_validation.json').write_text(json.dumps(out,ensure_ascii=False,indent=2));pd.DataFrame(ledgers).to_csv(O/'settlement_recalculation.csv',index=False)
    print(json.dumps(dict(status=out['status'],checks=len(checks),failures=[r for r in checks if r['status']=='FAIL']),ensure_ascii=False),flush=True)
    return out

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--workbooks',action='store_true');a=ap.parse_args();r=run(a.workbooks);raise SystemExit(0 if r['status']=='PASS' else 1)
