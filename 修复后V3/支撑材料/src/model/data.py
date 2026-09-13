from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook


def _date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value).strip().replace(".", "-").replace("/", "-"), "%Y-%m-%d").date()


def _hm(value: Any) -> str:
    if isinstance(value, time):
        return f"{value.hour:02d}:{value.minute:02d}"
    return str(value).strip()


def _numeric(values: list[Any], label: str) -> np.ndarray:
    bad = [i for i, v in enumerate(values) if v is None or v == "" or isinstance(v, bool)]
    if bad:
        raise ValueError(f"{label}: blank/non-numeric cells at offsets {bad[:10]}")
    try:
        arr = np.asarray(values, dtype=float)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label}: non-numeric cell present") from exc
    if not np.all(np.isfinite(arr)):
        raise ValueError(f"{label}: NaN/Inf present")
    return arr


@dataclass(frozen=True)
class DataBundle:
    dates: tuple[date, ...]
    source_timestamps: tuple[str, ...]
    official_labels: tuple[str, ...]
    fixed_price: np.ndarray
    q1_load_kw: np.ndarray
    q1_pv_kw: np.ndarray
    load_kw: np.ndarray
    pv_kw: np.ndarray
    realtime_price: np.ndarray
    pv_forecast_kw: dict[tuple[date, int], np.ndarray]

    @property
    def date_index(self) -> dict[date, int]:
        return {d: i for i, d in enumerate(self.dates)}

    def time_rows(self) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        for di, day in enumerate(self.dates):
            for pos in range(144):
                start = pos * 10
                end = (pos + 1) * 10
                rows.append({
                    "date": day.isoformat(),
                    "position": pos + 1,
                    "global_position": di * 144 + pos,
                    "source_timestamp": self.source_timestamps[pos],
                    "physical_interval": f"{start//60:02d}:{start%60:02d}-{end//60:02d}:{end%60:02d}",
                    "official_label": self.official_labels[pos],
                })
        return rows


def _read_matrix(path: Path, sheet: str) -> tuple[tuple[str, ...], tuple[date, ...], np.ndarray]:
    ws = load_workbook(path, data_only=True, read_only=True)[sheet]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    if len(header) < 145:
        raise ValueError(f"{path.name}/{sheet}: expected 145 columns")
    headers = tuple(_hm(x) for x in header[1:145])
    dates: list[date] = []
    matrix: list[np.ndarray] = []
    for ri, row in enumerate(rows, start=2):
        if row[0] in (None, ""):
            raise ValueError(f"{path.name}/{sheet}!A{ri}: blank date")
        dates.append(_date(row[0]))
        matrix.append(_numeric(list(row[1:145]), f"{path.name}/{sheet} row {ri}"))
    out = np.vstack(matrix)
    if out.shape != (365, 144) or len(set(dates)) != 365:
        raise ValueError(f"{path.name}/{sheet}: expected 365 unique days x 144 intervals, got {out.shape}")
    return headers, tuple(dates), out


def load_data(cfg: dict[str, Any]) -> DataBundle:
    root = Path(cfg["_source_root"]) / "C题" / "附件"
    ws1 = load_workbook(root / "附件1.xlsx", data_only=True, read_only=True).active
    rows1 = list(ws1.iter_rows(values_only=True))
    if len(rows1) != 145:
        raise ValueError("附件1: expected header + 144 intervals")
    source_timestamps = tuple(_hm(r[0]) for r in rows1[1:])
    fixed_price = _numeric([r[1] for r in rows1[1:]], "附件1/电价")
    q1_load = _numeric([r[2] for r in rows1[1:]], "附件1/负载")
    q1_pv = _numeric([r[3] for r in rows1[1:]], "附件1/光伏")

    h1, dates, load = _read_matrix(root / "附件2.xlsx", "小区负载")
    h2, dates2, pv = _read_matrix(root / "附件2.xlsx", "光伏发电实际功率")
    h3, dates3, rt = _read_matrix(root / "附件4.xlsx", "Sheet1")
    if not (dates == dates2 == dates3 and h1 == h2 == h3):
        raise ValueError("附件2/4 date or timestamp axes differ")
    expected = tuple(date(2025, 1, 1) + timedelta(days=i) for i in range(365))
    if dates != expected:
        raise ValueError("附件2/4 dates are not exactly 2025-01-01..2025-12-31")

    template = load_workbook(root / "附件5" / "result2.xlsx", data_only=True, read_only=True)["计划购电量"]
    official = tuple(str(x.value) for x in next(template.iter_rows())[1:145])
    if len(official) != 144 or len(set(official)) != 144:
        raise ValueError("official result template does not contain 144 unique interval labels")

    ws3 = load_workbook(root / "附件3.xlsx", data_only=True, read_only=True).active
    forecasts: dict[tuple[date, int], np.ndarray] = {}
    current: date | None = None
    for ri, row in enumerate(ws3.iter_rows(min_row=2, values_only=True), start=2):
        if row[0] not in (None, ""):
            current = _date(row[0])
        if current is None:
            raise ValueError(f"附件3 row {ri}: missing carried date")
        issue_text = str(row[1]).strip()
        try:
            issue = int(issue_text.split(":", 1)[0])
        except Exception as exc:
            raise ValueError(f"附件3 row {ri}: invalid issue time {row[1]!r}") from exc
        vals = _numeric(list(row[2:26]), f"附件3 row {ri}")
        key = (current, issue)
        if key in forecasts:
            raise ValueError(f"附件3 duplicate forecast {key}")
        forecasts[key] = vals
    expected_keys = {(d, h) for d in dates for h in (0, 6, 12, 18)}
    if set(forecasts) != expected_keys:
        raise ValueError(f"附件3 forecast coverage mismatch: missing={len(expected_keys-set(forecasts))}, extra={len(set(forecasts)-expected_keys)}")
    return DataBundle(dates, source_timestamps, official, fixed_price, q1_load, q1_pv, load, pv, rt, forecasts)
