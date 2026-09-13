"""Generate the documented AI-use description without inferring human sign-off."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]


def main():
    data=json.loads((ROOT/'evidence/ai_usage.json').read_text())
    doc=Document()
    for style in doc.styles:
        for border in list(style.element.xpath('.//w:pBdr')):border.getparent().remove(border)
    s=doc.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Cm(2.5)
    doc.add_paragraph(data['title'],'Title')
    for section in data['sections']:
        doc.add_paragraph(section['heading'],'Heading 1')
        for text in section['paragraphs']:doc.add_paragraph(text)
    for p in doc.paragraphs:
        title=p.style.name=='Title';heading=p.style.name=='Heading 1'
        f=p.paragraph_format
        f.alignment=WD_ALIGN_PARAGRAPH.CENTER if title or heading else WD_ALIGN_PARAGRAPH.JUSTIFY
        f.first_line_indent=Pt(0 if title or heading else 24)
        f.line_spacing=1;f.space_before=Pt(12 if heading else 0);f.space_after=Pt(12 if heading or title else 0)
        f.keep_with_next=title or heading;f.keep_together=True;f.widow_control=True
        for r in p.runs:
            r.font.name='Times New Roman';r.font.size=Pt(16 if title else 14 if heading else 12)
            r.font.bold=title or heading;r.font.color.rgb=RGBColor(0,0,0)
            fonts=r._r.get_or_add_rPr().get_or_add_rFonts();fonts.set(qn('w:eastAsia'),'SimSun')
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title=data['title']
    target=ROOT/'work/ai_usage/AI工具使用详情.docx';target.parent.mkdir(parents=True,exist_ok=True);doc.save(target)
    md='# '+data['title']+'\n\n'+''.join('## '+s['heading']+'\n\n'+'\n\n'.join(s['paragraphs'])+'\n\n' for s in data['sections'])
    (ROOT/'evidence/AI工具使用详情_冻结稿.md').write_text(md)
    print(target)


if __name__=='__main__':main()
