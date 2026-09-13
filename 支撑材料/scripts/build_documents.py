#!/usr/bin/env python3
"""Inject frozen results into the paper and create the final DOCX sources."""
from __future__ import annotations

import csv
import json
import re
import shutil
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
FINAL = ROOT / "outputs" / "c_final_v1"
FROZEN = FINAL / "frozen"
PAPER = FINAL / "paper"
CHARTS = FINAL / "charts"
WORK = ROOT / "work" / "paper_build"
# Arial Unicode MS is available as a TrueType font to the bundled headless
# converter and contains the complete Simplified Chinese glyph set. macOS TTC
# family names such as Songti SC/Heiti SC are visible to desktop applications
# but are not resolved reliably by this conversion runtime.
BODY_FONT = "Arial Unicode MS"
HEADING_FONT = "Arial Unicode MS"


def clean_inline(value: str) -> str:
    return value.replace('**', '').replace('`', '')


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr(); shd = tc_pr.find(qn("w:shd"))
    if shd is None: shd = OxmlElement("w:shd"); tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr(); borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None: borders = OxmlElement("w:tcBorders"); tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None: node = OxmlElement(f"w:{edge}"); borders.append(node)
        node.set(qn("w:val"), "single"); node.set(qn("w:sz"), "4"); node.set(qn("w:color"), color)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr(); tbl_header = OxmlElement("w:tblHeader"); tbl_header.set(qn("w:val"), "true"); tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(); begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def configure(doc: Document) -> None:
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    section = doc.sections[0]; section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5); section.bottom_margin = Cm(2.5); section.left_margin = Cm(2.5); section.right_margin = Cm(2.5)
    styles = doc.styles
    normal = styles["Normal"]; normal.font.name = BODY_FONT; normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.line_spacing = 1.25; normal.paragraph_format.space_after = Pt(3)
    for style_name, size, font in (("Title", 18, HEADING_FONT), ("Heading 1", 14, HEADING_FONT), ("Heading 2", 12, HEADING_FONT), ("Heading 3", 11, HEADING_FONT)):
        st = styles[style_name]; st.font.name = font; st.font.size = Pt(size); st.font.color.rgb = RGBColor(0, 0, 0); st.font.bold = True
        st._element.rPr.rFonts.set(qn("w:eastAsia"), font); st.paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.space_after = Pt(12)
    for st in styles:
        if st.type == 1:
            pr=st._element.find(qn('w:pPr'))
            if pr is not None:
                for border in list(pr.findall(qn('w:pBdr'))):pr.remove(border)
    table_style=styles.add_style('ResultTableText',1)
    table_style.font.name=BODY_FONT;table_style.font.size=Pt(8.5)
    table_style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),BODY_FONT)
    table_style.paragraph_format.line_spacing=1.0
    table_style.paragraph_format.space_before=Pt(0);table_style.paragraph_format.space_after=Pt(0)
    add_page_number(section.footer.paragraphs[0])


def add_table(doc: Document, lines: list[str]) -> None:
    data = [[clean_inline(x.strip()) for x in line.strip().strip("|").split("|")] for line in lines if not re.match(r"^\s*\|?\s*:?-+", line)]
    if not data: return
    cols = max(map(len, data)); table = doc.add_table(rows=len(data), cols=cols); table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, row in enumerate(data):
        for j in range(cols):
            cell = table.cell(i, j); cell.text = row[j] if j < len(row) else ""; cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER; set_cell_borders(cell)
            for p in cell.paragraphs:
                p.style=doc.styles['ResultTableText']
                p.paragraph_format.space_after = Pt(0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 or i == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = BODY_FONT; run.font.size = Pt(8.5); run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
                    if i == 0: run.font.bold = True; run.font.color.rgb = RGBColor(255,255,255)
            if i == 0: set_cell_shading(cell, "1F4E78")
            elif i % 2 == 0: set_cell_shading(cell, "F4F7FA")
    for row in table.rows:
        prevent_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(prevent_split)
        for j, cell in enumerate(row.cells):
            cell.width = Cm(16 / cols)
    if len(data) <= 12:
        for row in table.rows[:-1]:
            for cell in row.cells:
                for p in cell.paragraphs: p.paragraph_format.keep_with_next = True
    repeat_header(table.rows[0])
    spacer=doc.add_paragraph();spacer.paragraph_format.line_spacing=Pt(3)
    spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.space_after=Pt(4)


def markdown_to_docx(text: str, output: Path) -> None:
    doc = Document(); configure(doc)
    lines = text.splitlines(); i = 0; in_code = False; figure_no=0; table_no=0; current_heading=''
    while i < len(lines):
        line = lines[i]
        if line.startswith("``` ") or line.startswith("```"):
            in_code = not in_code; i += 1; continue
        if in_code:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(line if line else " ")
            run.font.name = "Consolas"; run.font.size = Pt(7.5)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
            i += 1; continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
            table_no+=1
            caption=re.sub(r'^\d+(?:\.\d+)*\s+', '', current_heading)
            columns=[x.strip() for x in line.strip('|').split('|')]
            names={'场景候选':'一月验证期候选策略费用','方案':'确定性调度与无储能对照','物理区间':'指定物理区间购电量' if '0点计划量' in columns else '指定物理区间储能动作','0点SOC':'日初与日末储能状态','全天计划量':'全天购电及费用汇总','紧急购电物理区间':'连续紧急购电区间'}
            caption=names.get(columns[0],caption)
            if columns[0]=='模式':
                caption={'方案':'日前采购与可行基线','最终净额':'同一承诺链的不同结算解释','一月启动费用':'一月启动与二月状态衔接','计划费':'正式期费用构成','时点策略':'更新频次的费用与紧急购电量','比较':'配对费用差及区间估计','紧急时段占比':'紧急购电与日费用尾部风险','变量':'执行时刻预测误差','求解方法':'全年替代求解器回放费用','块长/天':'分块重采样对块长度的敏感性'}.get(columns[1],caption)
            cap=doc.add_paragraph(f'表 {table_no}  {caption}')
            cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;cap.paragraph_format.keep_with_next=True
            for run in cap.runs:run.font.size=Pt(9)
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i]); i += 1
            add_table(doc, block); continue
        fig = re.fullmatch(r"\[\[FIGURE:([^|]+)\|(.+)\]\]", line.strip())
        if fig:
            figure_no+=1
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = False
            p.add_run().add_picture(str(CHARTS / fig.group(1)), width=Cm(15.2))
            p.add_run().add_break()
            cap=p;cap.add_run(f"图 {figure_no}  {fig.group(2)}"); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_together = True
            for run in cap.runs: run.font.size = Pt(9); run.font.name = BODY_FONT; run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
            i += 1; continue
        if line.startswith("# "):
            p = doc.add_paragraph(clean_inline(line[2:]), style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            heading = clean_inline(line[3:])
            current_heading=heading
            p=doc.add_paragraph(heading, style="Heading 1")
            if heading.startswith('1 ') or heading=='附录':p.paragraph_format.page_break_before=True
        elif line.startswith("### "):
            current_heading=clean_inline(line[4:])
            p=doc.add_paragraph(clean_inline(line[4:]), style="Heading 2")
            if re.match(r'### 附录[B-D] ',line):p.paragraph_format.page_break_before=True
        elif line.startswith("#### "):
            heading=clean_inline(line[5:])
            current_heading=heading
            p=doc.add_paragraph(heading, style="Heading 3")
        elif re.match(r"^\d+\.\s+", line):
            doc.add_paragraph(clean_inline(re.sub(r"^\d+\.\s+", "", line)), style="List Number")
        elif line.startswith("- "):
            doc.add_paragraph(clean_inline(line[2:]), style="List Bullet")
        elif line.startswith("$$") and line.endswith("$$"):
            from native_equations import equation
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            number=int(re.search(r'（(\d+)）',line).group(1));p._p.append(equation(number))
        elif not line.strip():
            pass
        else:
            p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Pt(21)
            p.paragraph_format.keep_together = False
            run = p.add_run(clean_inline(line))
            if re.match(r'^\[\d+\] ',line):
                p.paragraph_format.first_line_indent=Pt(0)
                p.paragraph_format.line_spacing=1.05
                p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.keep_together=True
                run.font.size=Pt(9)
        i += 1
    output.parent.mkdir(parents=True, exist_ok=True); doc.save(output)


def main():
    from paper_content import generate
    text,ai=generate()
    from research_content import enrich
    text,ai=enrich(text,ai)
    PAPER.mkdir(parents=True,exist_ok=True);WORK.mkdir(parents=True,exist_ok=True)
    (FINAL/'evidence/paper_source_frozen.md').write_text(text,encoding='utf-8')
    (FINAL/'evidence/AI工具使用详情_冻结稿.md').write_text(ai,encoding='utf-8')
    markdown_to_docx(text,PAPER/'C题论文.docx')
    from research_content import append_sources
    append_sources(PAPER/'C题论文.docx')
    markdown_to_docx(ai,PAPER/'AI工具使用详情.docx')
    print('Created current-run DOCX sources',flush=True)

if __name__=='__main__':main()
