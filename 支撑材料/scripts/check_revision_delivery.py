"""Structural and numerical publication gates; visual inspection is recorded separately."""
from pathlib import Path
import getpass,hashlib,json,re
from docx import Document
from pypdf import PdfReader
import pandas as pd
import openpyxl
ROOT=Path(__file__).resolve().parents[1]


def main():
    paper=ROOT/'publication/C题论文_实质改进版.docx';pdf=paper.with_suffix('.pdf')
    doc=Document(paper);reader=PdfReader(pdf)
    texts=[p.extract_text() for p in reader.pages]
    appendix=next(i+1 for i,t in enumerate(texts) if re.search(r'^附录$',t,re.M))
    ref=next(i+1 for i,t in enumerate(texts) if re.search(r'^参考文献$',t,re.M))
    assert '针对问题一' in texts[0] and '针对问题四' in texts[0] and '关键词' in texts[0]
    assert '问题重述' not in texts[0]
    assert appendix-2<=30,'Official main-body page limit exceeded'
    assert appendix-1<=25,'User requested at most 25 pages including abstract and references'
    assert all(abs(float(p.mediabox.width)-595.3)<1 and abs(float(p.mediabox.height)-841.9)<1 for p in reader.pages)
    for s in doc.sections:
        assert min(s.top_margin.cm,s.bottom_margin.cm,s.left_margin.cm,s.right_margin.cm)>=2.49
    for forbidden in [getpass.getuser(),'wx'+'id_']:
        assert forbidden not in '\n'.join(texts),f'Identity marker in PDF: {forbidden}'
    assert not doc.core_properties.author and not doc.core_properties.last_modified_by
    assert pdf.stat().st_size<20_000_000 and paper.stat().st_size<20_000_000
    assert all(len(re.sub(r'\s|\d','',t))>0 for t in texts),'Empty page'
    ps=[p.text for p in doc.paragraphs]
    records=json.loads((ROOT/'evidence/source_appendix_manifest_revision.json').read_text())
    for r in records:
        src=ROOT/r['path'];assert hashlib.sha256(src.read_bytes()).hexdigest()==r['sha256']
        lines=src.read_text().splitlines();i=ps.index(r['path'])
        assert ps[i+1]=='SHA256 '+r['sha256']
        assert ps[i+2:i+2+len(lines)]==lines,('Source block mismatch',r['path'])
    meta=json.loads((ROOT/'论文修订源文件/build_meta.json').read_text());checks=[]
    for key,num in meta['tables'].items():
        if not key.startswith('annex_'):continue
        table=doc.tables[num-1]
        if key.startswith('annex_storage'):
            assert len(table.columns)==6 and len(table.rows)==5
            mode=key[len('annex_storage'):].split('2025')[0];dt=re.search(r'2025-\d\d-\d\d',key).group()
            d=pd.read_csv(ROOT/f'outputs/c_final_v1/frozen/{mode}_dispatch.csv',float_precision='round_trip');z=d[d.date==dt]
            for ri,pair in enumerate([[0,4],[8,12],[16,20]],start=1):
                for ci,hour in enumerate(pair):
                    part=z[(z.position>hour*6)&(z.position<=(hour+4)*6)]
                    for offset,col in [(1,'charge_kwh'),(2,'discharge_kwh')]:
                        assert abs(float(table.rows[ri].cells[3*ci+offset].text)-part[col].sum())<=.0000501
            assert abs(float(table.rows[-1].cells[1].text)-z.soc_start_kwh.iloc[0])<=.0000501
            assert abs(float(table.rows[-1].cells[4].text)-z.soc_end_kwh.iloc[-1])<=.0000501
        elif key.startswith('annex_emergency'):
            assert len(table.columns)==8 and len(table.rows)>=3
            mode=key[len('annex_emergency'):]
            name={'q2':'result2','q3':'result3','q4_2':'result4-2','q4_3':'result4-3'}[mode]
            wb=openpyxl.load_workbook(ROOT/f'workbooks/{name}.xlsx',read_only=True,data_only=True)
            values={};day=None
            for row in list(wb['紧急购电量'].values)[1:]:
                if row[0] is not None:day=str(row[0])[:10]
                values.setdefault(day,[]).append((row[1],float(row[2])))
            for ci,dt in enumerate(['2025-03-20','2025-06-21','2025-09-23','2025-12-21']):
                actual=[(row.cells[2*ci].text,float(row.cells[2*ci+1].text)) for row in table.rows[2:] if row.cells[2*ci].text!='—']
                expected=values[dt]
                assert [x[0] for x in actual]==[x[0] for x in expected],('Emergency interval mismatch',mode,dt)
                assert max(abs(a[1]-b[1]) for a,b in zip(actual,expected))<=.000051,('Emergency rounding mismatch',mode,dt)
            wb.close()
        else:
            assert len(table.columns)==6 and len(table.rows)==4
            mode=key[len('annex_'):].split('2025')[0]
            dt=re.search(r'2025-\d\d-\d\d',key).group()
            col='plan_00_grid_kwh' if key.endswith('plan_00_grid_kwh') else 'final_adjusted_grid_kwh'
            d=pd.read_csv(ROOT/f'outputs/c_final_v1/frozen/{mode}_dispatch.csv',float_precision='round_trip')
            z=d[d.date==dt]
            for ri,hours in [(1,[10,12,14]),(2,[16,18,20])]:
                for ci,hour in enumerate(hours):
                    val=float(table.rows[ri].cells[2*ci+1].text)
                    expected=float(z[z.position==hour*6+1][col].iloc[0])
                    assert abs(val-expected)<=.0000501,(key,hour,val,expected)
        checks.append(key)
    statement=(ROOT/'论文修订源文件/ai_statement.txt').read_text().strip()
    assert statement in ps,'AI statement differs from current source'
    assert ps.index('AI工具使用声明')<ps.index('参考文献')
    listed=json.loads((ROOT/'evidence/support_file_list.json').read_text())['files']
    assert all(path in ps for path in listed),'Incomplete supporting-file appendix'
    report=dict(pdf_pages=len(reader.pages),abstract_page=1,reference_page=ref,appendix_starts_page=appendix,
        main_including_abstract_references=appendix-1,body_excluding_abstract=appendix-2,
        docx_bytes=paper.stat().st_size,pdf_bytes=pdf.stat().st_size,source_blocks=len(records),
        specified_horizontal_tables=len(checks),table_count=len(doc.tables),
        docx_sha256=hashlib.sha256(paper.read_bytes()).hexdigest(),pdf_sha256=hashlib.sha256(pdf.read_bytes()).hexdigest(),
        structural_and_numeric='PASS',visual_review='separate gate; this script does not certify visual quality')
    (ROOT/'publication/structural_checks.json').write_text(json.dumps(report,indent=2))
    print(json.dumps(report),flush=True)


if __name__=='__main__':main()
