"""Validate generated paper tables against dispatch positions, plus PDF and assets."""
from pathlib import Path
import json,re,hashlib
from datetime import datetime
import numpy as np,pandas as pd
from docx import Document
from pypdf import PdfReader
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];FINAL=ROOT/'outputs/c_final_v1';F=FINAL/'frozen'

def main():
    checks=[]
    def check(k,ok,value):checks.append(dict(check=k,status='PASS' if bool(ok) else 'FAIL',actual=value))
    audit=json.loads((FINAL/'audit/independent_validation.json').read_text());check('independent_model_and_workbooks',audit['status']=='PASS',audit['status'])
    doc=Document(FINAL/'paper/C题论文.docx');pdf=PdfReader(FINAL/'paper/C题论文.pdf');ai=PdfReader(FINAL/'paper/AI工具使用详情.pdf')
    text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells);pdftext='\n'.join(p.extract_text() for p in pdf.pages)
    # Literal template syntax inside the printed source appendix is source code,
    # not an unexpanded value in the paper. Check the rendered paper separately.
    body_paragraphs=[]
    for p in doc.paragraphs:
        if p.text=='附录C 完整源程序与运行配置':break
        body_paragraphs.append(p.text)
    source_page=next((i for i,p in enumerate(pdf.pages) if '完整源程序与' in p.extract_text()[:120]),len(pdf.pages))
    display_text='\n'.join(body_paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)+'\n'+''.join(p.extract_text() for p in pdf.pages[:source_page])
    check('no_unresolved_placeholders',not re.search(r'\{\{[^}]+\}\}',display_text),len(re.findall(r'\{\{[^}]+\}\}',display_text)))
    check('A4_pages',all(abs(float(p.mediabox.width)-595.276)<2 and abs(float(p.mediabox.height)-841.89)<2 for p in pdf.pages),len(pdf.pages))
    check('PDF_under_20MB',(FINAL/'paper/C题论文.pdf').stat().st_size<20*1024**2,(FINAL/'paper/C题论文.pdf').stat().st_size)
    # Walk actual OOXML order; attach each appendix table to its section date.
    active=None;seen={};errors=[];table_i=0;para_i=0
    frames={mode:pd.read_csv(F/f'{mode}_dispatch.csv') for mode in ['q1','q2','q3','q4_2','q4_3']}
    def num(v):return float(v.replace(',','').strip())
    for el in doc.element.body:
        if el.tag==qn('w:p'):
            t=''.join(el.itertext()) # itertext may repeat run text; regex still identifies section.
            m=re.search(r'(Q1|Q2|Q3|Q4_2|Q4_3) (2025-\d\d-\d\d) 指定结果',t)
            if m:active=(m.group(1).lower(),m.group(2));seen.setdefault(active,set())
        elif el.tag==qn('w:tbl'):
            table=doc.tables[table_i];table_i+=1
            if active is None:continue
            data=[[c.text for c in r.cells] for r in table.rows];head=data[0];mode,day=active;z=frames[mode];z=z[z.date==day]
            if head==['物理区间','0点计划量','最终承诺量']:
                seen[active].add('ten_minute')
                if len(data)!=7:errors.append((active,'ten_minute_count'))
                for row,h in zip(data[1:],[10,12,14,16,18,20]):
                    r=z[z.position==h*6+1].iloc[0]
                    if row[0]!=r.physical_interval or abs(num(row[1])-r.plan_00_grid_kwh)>.000051 or abs(num(row[2])-r.final_adjusted_grid_kwh)>.000051:errors.append((active,row))
            elif head==['物理区间','充电量','放电量']:
                seen[active].add('four_hour')
                for b,row in enumerate(data[1:]):
                    zz=z[(z.position>b*24)&(z.position<=(b+1)*24)]
                    if row[0]!=f'{b*4}:00-{(b+1)*4}:00' or abs(num(row[1])-zz.charge_kwh.sum())>.000051 or abs(num(row[2])-zz.discharge_kwh.sum())>.000051:errors.append((active,row))
            elif head==['0点SOC','24点SOC']:
                seen[active].add('soc')
                if abs(num(data[1][0])-z.soc_start_kwh.iloc[0])>.000051 or abs(num(data[1][1])-z.soc_end_kwh.iloc[-1])>.000051:errors.append((active,'SOC'))
            elif head==['全天计划量','全天最终量','全天总购电费']:
                seen[active].add('daily')
                if abs(num(data[1][0])-z.plan_00_grid_kwh.sum())>.000051 or abs(num(data[1][1])-z.final_adjusted_grid_kwh.sum())>.000051 or abs(num(data[1][2])-z.final_relative_cost_yuan.sum())>.0051:errors.append((active,'daily'))
            elif head==['紧急购电物理区间','紧急购电量']:
                seen[active].add('emergency');expected=[];start=None;total=0.;last=None
                for r in z.sort_values('position').itertuples():
                    if r.emergency_kwh>=.00005:
                        if start is None:start=(int(r.position)-1)*10
                        last=int(r.position)*10;total+=r.emergency_kwh
                    elif start is not None:
                        expected.append((f'{start//60}:{start%60:02d}-{last//60}:{last%60:02d}',total));start=None;total=0.
                if start is not None:expected.append((f'{start//60}:{start%60:02d}-{last//60}:{last%60:02d}',total))
                if not expected:expected=[('无正缺口',0.)]
                if len(data)-1!=len(expected):errors.append((active,'emergency_count'))
                for row,(span,value) in zip(data[1:],expected):
                    if row[0]!=span or abs(num(row[1])-value)>.000051:errors.append((active,'emergency',row))
    required={(mode,day) for mode in ['q2','q3','q4_2','q4_3'] for day in ['2025-03-20','2025-06-21','2025-09-23','2025-12-21']};required.add(('q1','2025-01-01'))
    missing=[k for k in required if k not in seen or not ({'ten_minute','four_hour','soc','daily'}|({'emergency'} if k[0]!='q1' else set()))<=seen[k]]
    check('all_17_specified_date_groups',not missing,{'groups':len(seen),'missing':missing});check('paper_numbers_independently_reconciled',not errors,errors)
    missing_pdf_values=[]
    compact_pdf=re.sub(r'\s+','',pdftext)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                value=cell.text.strip()
                if re.fullmatch(r'-?[\d,]+\.\d+',value) and value not in compact_pdf:missing_pdf_values.append(value)
    check('all_table_numeric_strings_survive_PDF',not missing_pdf_values,missing_pdf_values)
    # Confirm key corrected Q1 numbers survived the PDF conversion.
    from decimal import Decimal,ROUND_HALF_UP
    q1value=float(frames['q1'].loc[frames['q1'].position==73,'plan_00_grid_kwh'].iloc[0])
    for value in [format(Decimal(str(q1value)).quantize(Decimal('.0001'),rounding=ROUND_HALF_UP),'.4f')]:
        check('PDF_corrected_physical_value_'+value,value in pdftext,value)
    check('five_workbooks',len(list((FINAL/'workbooks').glob('result*.xlsx')))==5,len(list((FINAL/'workbooks').glob('result*.xlsx'))))
    check('AI_details_nonempty',len(ai.pages)>0 and len(''.join(p.extract_text() for p in ai.pages))>200,len(ai.pages))
    check('charts_present',len(list((FINAL/'charts').glob('*.png')))>=4,len(list((FINAL/'charts').glob('*.png'))))
    (FINAL/'qa/paper_text.txt').write_text(pdftext)
    from competition_checks import inspect
    checks.extend(inspect(doc,pdf))
    result=dict(status='PASS' if all(c['status']=='PASS' for c in checks) else 'FAIL',generated_at=datetime.now().astimezone().isoformat(),scope='current_run_numeric_document_and_export_checks',checks=checks,metrics=dict(pdf_pages=len(pdf.pages),docx_tables=len(doc.tables),specified_date_groups=len(seen)),limitations=['Historical scenarios do not guarantee unseen tail-risk coverage.','Refund and delivery-time pricing are declared interpretations, not additional official rules.','Solver diagnostics cover 32 exact LP instances, not full-year alternative-solver policy rollouts.','Team submission approval and authorship signoff are not asserted.'])
    result['limitations']=['Finite historical scenarios and one-year replays do not guarantee future tail-risk coverage.','Refund and delivery-time pricing remain declared interpretations.','Annual alternative-solver checks cover S2 and S3 with a fixed common reference contract; no test-period policy selection.','Team authorship and manual review are pending; this PASS is not a submission eligibility certificate.']
    result['submission_readiness']='PENDING_TEAM_MANUAL_REVIEW_AND_RULES_CONFIRMATION'
    artifacts=list((FINAL/'paper').glob('*.docx'))+list((FINAL/'paper').glob('*.pdf'))+list((FINAL/'workbooks').glob('*.xlsx'))
    result['artifact_hashes']=[dict(path=str(p.relative_to(FINAL)),sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in artifacts]
    (FINAL/'final_validation_v3.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False,indent=2));return result

if __name__=='__main__':raise SystemExit(0 if main()['status']=='PASS' else 1)
