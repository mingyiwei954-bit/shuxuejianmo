from pathlib import Path
from copy import deepcopy
import re,json,hashlib,sys
import pandas as pd
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
W=Path(__file__).resolve().parent
SUPPORT=W/'support' if (W/'support').exists() else W.parent
sys.path.insert(0,str(SUPPORT/'scripts'))
import native_equations as ne
old=Document(W/'source/current_v3.docx');d=Document(W/'source/current_v3.docx');ps=old.paragraphs
for x in list(d.element.body):
 if x.tag!=qn('w:sectPr'):d.element.body.remove(x)
nums={};figs={};eqs={};seen_cites=[];table_items=[];caption_ids=set();code_ids=set();main_paras=[]
def p(text='',style=None):
 x=d.add_paragraph(text,style);main_paras.append(x);return x
def h(text,level):return p(text,'Heading '+str(level))
def push(node):
 el=deepcopy(node);d.element.body.insert(len(d.element.body)-1,el);return el
def oldp(i):
 x=p(ps[i].text);return x
def table(key,title=None,rows=None,widths=None):
 num=len(nums)+1;nums[str(key)]=num
 if isinstance(key,int):
  cap=next(x.text for x in ps if re.match(r'^表\s*'+str(key)+r'\s',x.text));title=re.sub(r'^表\s*\d+\s*','',cap)
 if key==1:title='主要符号及单位'
 c=p(f'表 {num}  {title}');caption_ids.add(c._p)
 if rows is None:
  from docx.table import Table
  t=Table(push(old.tables[key-1]._tbl),d)
 else:
  t=d.add_table(rows=1,cols=len(rows[0]))
  for cell,s in zip(t.rows[0].cells,rows[0]):cell.text=str(s)
  for vals in rows[1:]:
   for cell,s in zip(t.add_row().cells,vals):cell.text=str(s)
 table_items.append((str(key),t,widths));return t
oldfig={int(re.search(r'图\s*(\d+)',x.text).group(1)):x for x in ps if x._p.xpath('.//w:drawing')}
def fig(key):
 n=len(figs)+1;figs[str(key)]=n
 if isinstance(key,int):
  el=push(oldfig[key]._p);x=Paragraph(el,d)
  for r in x.runs:
   if re.search(r'图\s*\d+',r.text):r.text=re.sub(r'图\s*\d+',f'图 {n}',r.text)
 else:
  title={'route':'微网购电调度与检验技术路线','sensitivity':'固定窗口对终端惩罚与功率参数的响应'}[key]
  x=p();x.add_run().add_picture(str(W/'figures'/f'{key}.png'),width=Cm(15.8));x.add_run(f'\n图 {n}  {title}')
 caption_ids.add(x._p);return x
oldmath={int(re.search(r'\((\d+)\)',x.text).group(1)):x for x in ps if x._p.xpath('.//m:oMath')}
def equation(key):
 n=len(eqs)+1;eqs[str(key)]=n
 if isinstance(key,int):
  x=Paragraph(push(oldmath[key]._p),d)
  for r in x.runs:
   if re.search(r'\(\d+\)',r.text):r.text=re.sub(r'\(\d+\)',f'({n})',r.text)
 else:
  texts={'q1':r'min J_1 = ∑_{t=1}^{144} p_t B_t； E_0=E_{144}=6000',
         'price':r'E[p_t Q_t] = E[p_t] E[Q_t] + Cov(p_t,Q_t)',
         'residual':r'r_t^E = E_{t+1}−E_t−η_c C_t+\frac{D_t}{η_d}'}
  x=p();x.add_run('\t');x._p.append(ne.node('oMath',ne.parse(texts[key])));x.add_run(f'\t({n})')
 return x
refs={
1:'全国大学生数学建模竞赛组委会. 2026年高教社杯全国大学生数学建模竞赛C题：微网与外部电网电力调控策略及附件[Z]. 2026.',
3:'SHAPIRO A, DENTCHEVA D, RUSZCZYNSKI A. Lectures on stochastic programming: modeling and theory[M]. 3rd ed. Philadelphia: SIAM, 2021. DOI:10.1137/1.9781611976595.',
4:'ALARCON M A, ALARCON R G, GONZALEZ A H, et al. A scenario-based economic-stochastic model predictive control for the management of microgrids[J]. Sustainable Energy, Grids and Networks, 2023, 36: 101205. DOI:10.1016/j.segan.2023.101205.',
5:'HYNDMAN R J, ATHANASOPOULOS G. 预测：方法与实践[M/OL]. 3版. Melbourne: OTexts, 2021[2026-09-13]. https://otexts.com/fpp3cn/.',
6:'KUNSCH H R. The jackknife and the bootstrap for general stationary observations[J]. The Annals of Statistics, 1989, 17(3): 1217-1241. DOI:10.1214/aos/1176347265.',
7:'HUANGFU Q, HALL J A J. Parallelizing the dual revised simplex method[J]. Mathematical Programming Computation, 2018, 10(1): 119-142. DOI:10.1007/s12532-017-0130-5.',
8:'JOSHAL K S, GUPTA N. Microgrids with model predictive control: a critical review[J]. Energies, 2023, 16(13): 4851. DOI:10.3390/en16134851.',
9:'LIN Y, LI L, ZHANG J, et al. A scenario-based stochastic model predictive control approach for microgrid operation at an Australian cotton farm under uncertainties[J]. International Journal of Electrical Power & Energy Systems, 2024, 159: 110025. DOI:10.1016/j.ijepes.2024.110025.',
10:'BABIC L, LAURICELLA M, CEUSTERS G, et al. Data-driven non-parametric chance-constrained model predictive control for microgrids energy management using small data batches[J]. Frontiers in Control Engineering, 2023, 4: 1237759. DOI:10.3389/fcteg.2023.1237759.'}
p('基于场景线性规划的微网滚动购电调度','Title');h('摘要',1)
abstract=(W/'abstract.txt').read_text().splitlines()
for s in abstract:p(s)
p('关键词：场景随机规划 HiGHS算法 微网购电 滚动优化 参数扰动检验')
for line in (W/'body.md').read_text().splitlines():
 if not line:continue
 if line.startswith('#'):
  level=len(line)-len(line.lstrip('#'));x=h(line[level+1:],level)
  if line=='# 1 问题重述':x.paragraph_format.page_break_before=True
 elif line.startswith('@P '):
  for i in map(int,line[3:].split(',')):oldp(i)
 elif line.startswith('@E '):equation(int(line[3:]))
 elif line.startswith('@EQ '):equation(line[4:])
 elif line.startswith('@FIG '):
  k=line[5:];fig(int(k) if k.isdigit() else k)
 elif line.startswith('@T '):
  k=line[3:]
  if k.isdigit():
   t=table(int(k))
   if k=='1':
    t.rows[7].cells[1].text='窗口集合、场景索引与场景数量'
    for vals in [('Fτ，τ','决策信息集与发布时间','—，h'),('k，ΔAₜᵏ','调整次数索引与承诺变化量','—，kWh'),('J，fₜ','总费用与调整费用函数','元'),('rₜᴱ','储电量递推残差','kWh')]:
     for c,s in zip(t.add_row().cells,vals):c.text=s
  elif k=='input':
   a=json.loads((SUPPORT/'research/parameter_perturbation/input_quality.json').read_text());rows=[['输入变量','数值数量','缺失/负值','范围']]
   for v in a:rows.append([v['variable']+('/(元/kWh)' if v['variable']=='实时电价' else '/kW'),str(v['values']),f"{v['missing']} / {v['negative']}",f"{v['minimum']:.4f}—{v['maximum']:.4f}"])
   table(k,'原始输入的完整性与范围检查',rows,[4,3,3,6])
  elif k=='params':
   table(k,'模型求解与实验配置',[
    ['项目','设置','用途'],['求解环境','Python 3.12.14；SciPy 1.16.1','线性规划'],['常用数据依赖','NumPy 2.3.5；pandas 2.2.3','矩阵与记录'],['区间/窗口','10分钟 / 24小时','离散与滚动'],['场景数/终端倍数','20 / 1.0','一月选择后锁定'],['原始/对偶可行容差','均为10⁻⁸','可行性检查'],['内点法最优容差','10⁻⁹','替代求解'],['吞吐量正则系数','10⁻⁷元/kWh','数值选择'],['统计重采样','每组2000次','区间估计']], [4.5,6,5.5])
  elif k=='sensitivity':
   data=pd.read_csv(SUPPORT/'research/parameter_perturbation/summary.csv');rows=[['参数','扰动/%','目标变化最大绝对值/%','采购变化最大L1/kWh','储电量最大变化/kWh']]
   for r in data.itertuples():
    if r.ratio==1:continue
    rows.append(['终端倍数' if r.parameter=='terminal_multiplier' else '额定功率',f'{(r.ratio-1)*100:+.0f}',f'{r.max_abs_objective_change_pct:.2f}',f'{r.max_grid_l1_kwh:.2f}',f'{r.max_soc_change_kwh:.2f}'])
   table(k,'两项参数梯度扰动的固定窗口结果',rows,[2.4,1.8,4.0,4.0,3.8])
 elif line=='@SENSITIVITY_TEXT':
  p('终端惩罚倍数上下扰动20%时，32个窗口的目标变化最大绝对值为1.83%，但采购轨迹累计绝对变化最大达5333.33千瓦时，储电量最大变化达4800千瓦时。额定功率降低20%时，目标变化最大达15.23%，超过5%；提高20%时最大变化约0.73%。这说明当前模型对功率下降更敏感，不能概括为所有参数扰动下均稳定。')
 elif line=='@ALGORITHM':
  x=p('算法1  基于共同日前合同的滚动调度');x.style=d.styles['Heading 3']
  for s in ['输入：已选配置、当前储电量、允许更新时点、可用历史及发布预报。','1  用仅零点更新的参考策略形成各日共同日前合同B。','2  按时间顺序读取下一允许更新时刻τ。','3  截取τ以前完整历史与截至τ已发布的预报，生成联合场景。','4  保持已执行区间不变；零点锁定共同B，后续允许调整剩余承诺。','5  求解共享采购与储能动作的线性规划，执行互斥变换与残差检查。','6  保存本日剩余区间的承诺变化；次日预览不写入正式合同。','7  仅执行至下一更新时点，以随后实际观测计算缺口、费用及状态。','8  未到评价终点则返回步骤2；年末施加明确的末状态约束。','输出：逐时段采购、充放电、储电量、紧急缺口、费用与完整承诺链。']:
   x=p(s);code_ids.add(x._p)
 else:p(line)
# Main text citations are reordered by first appearance and later superscripted.
for x in d.paragraphs:
 for m in re.finditer(r'\[(\d+)\]',x.text):
  k=int(m.group(1))
  if k in refs and k not in seen_cites:seen_cites.append(k)
assert len(seen_cites)==9,seen_cites
refmap={k:i+1 for i,k in enumerate(seen_cites)}
h('AI工具使用声明',1);p(ps[115].text)
h('参考文献',1).paragraph_format.page_break_before=True
reference_ids=set()
for oldnum in seen_cites:
 x=p(f'[{refmap[oldnum]}] '+refs[oldnum]);reference_ids.add(x._p)
# Append the existing annexes and exact source text, adding one supplemental source.
start=next(x for x in ps if x.text=='附录')
annex_nodes=[];active=False
for el in old.element.body:
 if el is start._p:active=True;continue
 if active and el.tag!=qn('w:sectPr'):annex_nodes.append(el)
annex_start=None;supp_added=False;appendix_code=False
for el in annex_nodes:
 if el.tag==qn('w:p'):
  text=''.join(el.xpath('.//w:t/text()'))
  if text.startswith('附录3 ') and not supp_added:
   h('2.2 参数扰动补充程序',2)
   src=SUPPORT/'scripts/parameter_perturbation.py'
   x=p('scripts/parameter_perturbation.py');x.style=d.styles['Heading 3']
   p('SHA256 '+hashlib.sha256(src.read_bytes()).hexdigest())
   for line in src.read_text().splitlines():
    x=p(line);code_ids.add(x._p)
   supp_added=True;appendix_code=False
  if text=='config/final.yaml':appendix_code=True
  new=push(el);x=Paragraph(new,d)
  if annex_start is None:annex_start=x
  if text.startswith('附录3 '):x.paragraph_format.page_break_before=True
  if appendix_code and not x.style.name.startswith('Heading'):code_ids.add(new)
  mat=re.match(r'^表\s*(\d+)\s+(.*)',x.text)
  if mat:
   key=mat.group(1);nums[key]=len(nums)+1;x.text=f'表 {nums[key]}  {mat.group(2)}';caption_ids.add(new)
 else:
  from docx.table import Table
  new=push(el)
  if el.tag==qn('w:tbl'):table_items.append(('annex',Table(new,d),None))
# Keep appendix separate from the 25-page manuscript budget.
annex_start.paragraph_format.page_break_before=True
# Add the supplemental output locations to the appendix narrative only.
for x in d.paragraphs:
 if x.text.startswith('参数候选、选择规则与费用见正文'):
  x.text+=' 新增固定窗口扰动的输入、参数、320条实例记录与汇总见research/parameter_perturbation，复算程序全文见附录2.2；该诊断未改动正式模型参数。'
 if x.text.startswith('运行入口为支撑材料/'):
  x.text=x.text.replace('完整版本见evidence/environment_versions.json。','完整版本见evidence/environment_versions.json。本次另附参数扰动程序及诊断记录。')
# Update reproduction entry points without changing any archived source blocks.
for x in d.paragraphs:
 if x.text.startswith('运行入口为支撑材料/'):
  x.text='数值流水线入口为支撑材料/run_pipeline.sh；本次论文排版另用论文修订源文件/build.py，旧流水线的文档步骤仍生成旧版。原始附件路径在config/final.yaml中配置，也可用CUMCM_SOURCE_ROOT覆盖。数据校验、模型训练、正式期运行与独立审查依序执行，任一步失败即停止。记录环境为Python 3.12.14、NumPy 2.3.5、SciPy 1.16.1、pandas 2.2.3及openpyxl 3.1.5，版本明细见evidence/environment_versions.json。配套包中“原始附件”保存输入Excel；frozen/columns.tar.xz保存完整冻结明细，用scripts/unpack_frozen.py恢复。年度与训练费用分开，三种结算口径分栏；中间统计保存在evidence及research目录，原文件与完整精度值均保留。'
 if x.text.startswith('支撑材料的必要文件按功能列示如下'):
  x.text='支撑材料按功能列示如下。原模型与审查源码在附录2.1给出，新增参数扰动程序见附录2.2；排版源文件另包交付。完整文件校验清单见包内SHA256SUMS.json。'
 if x.text.startswith('以下逐文件收录当前工程自编的全部'):
  x.text='以下逐文件收录原工程32个程序与配置文件。第三方依赖按requirements.txt安装；工作簿导出还需Node与@oai/artifact-tool。程序以支撑材料内原始文件运行，版面换行不改变源码；独立数值复验仅需Python。本次论文生成源文件与运行方法见论文修订源文件目录。'
 if x.text.startswith('以四个指定日期、两种电价、每天四次决策构成'):
  x.text=x.text.replace('全年策略费用的比较另见第7.3节。','全年策略费用采用下述连续回放另行比较。')
 if x.text.startswith('共完成96次固定实例复求'):
  x.text=x.text.replace('其全年影响进一步通过第7.3节的连续回放检验。','其全年影响进一步通过下述连续回放检验。')
# Apply number mappings in explanatory text, preserving code text and inline equations.
for x in d.paragraphs:
 if x._p in code_ids or x._p.xpath('.//m:oMath|.//w:drawing') or x._p in caption_ids or x._p in reference_ids:continue
 s=x.text
 s=re.sub(r'表\s*(\d+)',lambda m:'表'+str(nums.get(m.group(1),m.group(1))),s)
 s=re.sub(r'图\s*(\d+)',lambda m:'图'+str(figs.get(m.group(1),m.group(1))),s)
 s=re.sub(r'式（(\d+)）',lambda m:'式（'+str(eqs.get(m.group(1),m.group(1)))+'）',s)
 s=re.sub(r'@F(\d+)@',lambda m:str(figs[m.group(1)]),s)
 # Old structural references must match the rewritten outline.
 s=s.replace('第5节调整费用','第5.4节调整费用').replace('第7.3节','第6.4节')
 if s.startswith('参数候选、选择规则与费用见正文'):
  s='参数候选、选择规则与费用见正文表4；参数验证明细、解析后的锁定配置和原始输入校验记录随支撑材料交付。指定日期求解器复求、全年替代回放、月度费用、尾部风险、预测误差及块长度敏感性保存于research目录，完整文件路径与复算命令见配套阅读说明。新增参数扰动子目录保存输入检查、实验协议、320条实例及汇总，复算程序全文见附录2.2；这项诊断未改动正式模型参数和结果工作簿。'
 if s!=x.text:x.text=s
# Font and paragraph formatting follows the supplied requirements.
def rf(r,size=12,bold=False):
 r.font.name='Times New Roman';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
 pr=r._r.get_or_add_rPr()
 for a in list(pr.findall(qn('w:rFonts'))):pr.remove(a)
 e=OxmlElement('w:rFonts')
 for k,v in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','SimSun'),('cs','Times New Roman')]:e.set(qn('w:'+k),v)
 pr.insert(0,e)
def fp(x,size=12,bold=False,align=WD_ALIGN_PARAGRAPH.JUSTIFY,indent=24,code=False):
 f=x.paragraph_format;f.alignment=align;f.line_spacing=1;f.space_before=Pt(0);f.space_after=Pt(0);f.first_line_indent=Pt(indent);f.left_indent=Pt(0);f.right_indent=Pt(0);f.keep_together=not code;f.keep_with_next=False;f.widow_control=True
 for r in x.runs:rf(r,size,bold)
for x in d.paragraphs:
 iscode=x._p in code_ids;ishead=x.style.name.startswith('Heading');istitle=x.style.name=='Title';iscap=x._p in caption_ids;ismath=bool(x._p.xpath('.//m:oMath'))
 size=16 if istitle else 14 if x.style.name=='Heading 1' else 10.5 if iscap else 12
 align=WD_ALIGN_PARAGRAPH.CENTER if istitle or x.style.name=='Heading 1' or iscap else WD_ALIGN_PARAGRAPH.LEFT if ishead or iscode else WD_ALIGN_PARAGRAPH.JUSTIFY
 fp(x,size,ishead or istitle,align,0 if ishead or istitle or iscap or ismath or iscode else 24,iscode)
 if ishead or istitle:x.paragraph_format.keep_with_next=True;x.paragraph_format.space_before=Pt(12 if ishead else 0);x.paragraph_format.space_after=Pt(12)
 if iscap:
  x.paragraph_format.space_before=Pt(6);x.paragraph_format.space_after=Pt(3);x.paragraph_format.keep_with_next=not bool(x._p.xpath('.//w:drawing'))
 if ismath:
  x.alignment=WD_ALIGN_PARAGRAPH.LEFT;x.paragraph_format.tab_stops.clear_all();x.paragraph_format.tab_stops.add_tab_stop(Cm(8),WD_TAB_ALIGNMENT.CENTER);x.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT)
  for mr in x._p.xpath('.//m:r'):
   pr=mr.find(qn('w:rPr'))
   if pr is None:pr=OxmlElement('w:rPr');mr.insert(0,pr)
   for e in list(pr):pr.remove(e)
   e=OxmlElement('w:rFonts')
   for k,v in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','SimSun')]:e.set(qn('w:'+k),v)
   pr.append(e);e=OxmlElement('w:sz');e.set(qn('w:val'),'24');pr.append(e)
 if x._p in reference_ids:x.alignment=WD_ALIGN_PARAGRAPH.LEFT;x.paragraph_format.first_line_indent=Pt(0)
 elif not iscode and not ismath and not iscap and re.search(r'\[(\d+)\]',x.text):
  text=x.text;x.text=''
  for s in re.split(r'(\[\d+\])',text):
   if re.fullmatch(r'\[\d+\]',s):
    k=int(s[1:-1]);r=x.add_run('['+str(refmap.get(k,k))+']');rf(r);r.font.superscript=True
   else:rf(x.add_run(s))
# All table bodies use 12pt and three horizontal rules, preserving original numeric strings.
for key,t,widths in table_items:
 t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.style=d.styles['Normal Table'];pr=t._tbl.tblPr
 if widths:
  for col,v in zip(t.columns,widths):col.width=Cm(v)
 else:widths=[c.width.cm if c.width else 16/len(t.columns) for c in t.columns]
 for tag in ['w:tblBorders','w:shd']:
  for e in pr.findall(qn(tag)):pr.remove(e)
 bd=OxmlElement('w:tblBorders')
 for edge in ['top','bottom','left','right','insideH','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil');e.set(qn('w:sz'),'10');e.set(qn('w:color'),'000000');bd.append(e)
 pr.append(bd)
 for ri,row in enumerate(t.rows):
  rp=row._tr.get_or_add_trPr()
  for e in rp.findall(qn('w:trHeight')):rp.remove(e)
  if rp.find(qn('w:cantSplit')) is None:rp.append(OxmlElement('w:cantSplit'))
  for ci,c in enumerate(row.cells):
   c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cp=c._tc.get_or_add_tcPr()
   for tag in ['w:shd','w:tcBorders']:
    for e in cp.findall(qn(tag)):cp.remove(e)
   if ri==0:
    bd=OxmlElement('w:tcBorders');e=OxmlElement('w:bottom');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');bd.append(e);cp.append(bd)
   for x in c.paragraphs:fp(x,align=WD_ALIGN_PARAGRAPH.CENTER,indent=0);x.paragraph_format.keep_with_next=ri<len(t.rows)-1
for sec in d.sections:
 sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2.5)
d.core_properties.title=d.paragraphs[0].text;d.core_properties.author='';d.core_properties.last_modified_by=''
out=W/'C题论文_v3.docx';d.save(out)
# Basic content invariants before rendering.
assert 800<=sum(len(s) for s in abstract)<=1000
assert len(d.paragraphs[0].text)<=25
assert ps[115].text in [x.text for x in d.paragraphs]
origcode=[x.text for x in ps[ps.index(next(x for x in ps if x.text=='config/final.yaml')):]]
# Compare all 32 original source blocks until the next appendix.
a=next(i for i,x in enumerate(ps) if x.text=='config/final.yaml');z=next(i for i,x in enumerate(ps) if x.text.startswith('附录3 '))
oldseq=[x.text for x in ps[a:z]];newtexts=[x.text for x in d.paragraphs];a2=newtexts.index('config/final.yaml');assert newtexts[a2:a2+len(oldseq)]==oldseq
(W/'build_meta.json').write_text(json.dumps({'tables':nums,'figures':figs,'equations':eqs,'references':refmap,'abstract_chars':sum(map(len,abstract)),'abstract_hanzi':sum(len(re.findall('[\u4e00-\u9fff]',s)) for s in abstract),'table_count':len(d.tables),'paragraph_count':len(d.paragraphs)},ensure_ascii=False,indent=2))
print(out,'tables',len(d.tables),'figures',len(d.inline_shapes),'equations',len(eqs))
